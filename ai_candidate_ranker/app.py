#!/usr/bin/env python3
"""
Redrob Candidate Ranker Web Server - Role-Based Discovery
Flask-based backend for role-specific candidate discovery and ranking.
"""

import os
import json
import csv
from datetime import datetime
from flask import Flask, jsonify, request, send_from_directory, send_file
import io
from flask_cors import CORS

app = Flask(__name__)
CORS(app)

# Paths & Arguments Resolution
import sys
import argparse

parser = argparse.ArgumentParser(description="Redrob Candidate Discovery Server")
parser.add_argument("--candidates", type=str, default=None, help="Path to candidates.jsonl file")
parser.add_argument("--out", type=str, default=None, help="Path to output submission.csv")
# parse_known_args prevents crash when Flask runner passes options
args, unknown = parser.parse_known_args()

if getattr(sys, 'frozen', False):
    exe_dir = os.path.dirname(sys.executable)
else:
    exe_dir = os.path.dirname(os.path.abspath(__file__))

# Default search paths
CANDIDATES_PATH = os.path.join(exe_dir, "candidates.jsonl")
SUBMISSION_PATH = os.path.join(exe_dir, "team_antigravity.csv")

# Fallback checking
if args.candidates and os.path.exists(args.candidates):
    CANDIDATES_PATH = args.candidates
elif not os.path.exists(CANDIDATES_PATH):
    fallback_path = r"c:\Users\rajku\Downloads\[PUB] India_runs_data_and_ai_challenge\[PUB] India_runs_data_and_ai_challenge\India_runs_data_and_ai_challenge\candidates.jsonl"
    if os.path.exists(fallback_path):
        CANDIDATES_PATH = fallback_path

if args.out:
    SUBMISSION_PATH = args.out

# Global data structures
ALL_CANDIDATES = []
VALID_CANDIDATES = []
HONEYPOT_COUNT = 0
CURRENT_RANKING = []
CANDIDATE_MAP = {}
DATASET_STATS = {}
ACTIVE_ROLE = "senior_ai_engineer"

# Fictional companies list
FICTIONAL_COMPANIES = {"hooli", "stark industries", "wayne enterprises", "initech", "globex inc", "dunder mifflin"}

# Role Configurations
ROLE_CONFIGS = {
    "senior_ai_engineer": {
        "title": "Senior AI Engineer - Founding Team",
        "must_have_skills": {
            "embeddings", "retrieval", "vector search", "hybrid search", "pinecone",
            "weaviate", "qdrant", "milvus", "opensearch", "elasticsearch", "faiss",
            "sentence-transformers", "bge", "e5", "pytorch", "tensorflow", "nlp",
            "natural language processing", "information retrieval", "rag",
            "evaluation frameworks", "ndcg", "mrr", "map"
        },
        "nice_to_have_skills": {
            "llm", "llms", "fine-tuning", "lora", "qlora", "peft", "xgboost",
            "learning to rank", "learning-to-rank", "distributed systems", "inference optimization"
        },
        "irrelevant_skills": {
            "computer vision", "image classification", "object detection", "speech recognition",
            "robotics", "voice assistant", "gans", "yolo"
        },
        "title_keywords": ["ai engineer", "ml engineer", "machine learning engineer", "nlp engineer", "search engineer", "retrieval engineer", "data scientist", "applied scientist", "applied ml", "machine learning"],
        "min_yoe": 5,
        "max_yoe": 9,
        "optimal_yoe_low": 6,
        "optimal_yoe_high": 8,
        "industry_keywords": ["software", "fintech", "e-commerce"],
        "description_keywords": ["embeddings", "rag", "vector search", "pinecone", "recommender", "retrieval", "ndcg", "ranking"]
    },
    "applied_ml_engineer": {
        "title": "Applied Machine Learning Engineer",
        "must_have_skills": {
            "pytorch", "tensorflow", "scikit-learn", "numpy", "pandas", "xgboost",
            "deep learning", "machine learning", "feature engineering", "model deployment",
            "statistical modeling", "scipy", "keras", "regression", "classification"
        },
        "nice_to_have_skills": {
            "mlops", "docker", "kubernetes", "fastapi", "mlflow", "kubeflow", "bentoml",
            "data pipelines", "spark", "aws", "gcp"
        },
        "irrelevant_skills": {
            "marketing", "hr", "sales", "design", "figma", "photoshop"
        },
        "title_keywords": ["ml engineer", "machine learning engineer", "applied ml", "ml", "applied machine learning", "data scientist", "data science"],
        "min_yoe": 4,
        "max_yoe": 8,
        "optimal_yoe_low": 5,
        "optimal_yoe_high": 7,
        "industry_keywords": ["software", "fintech", "e-commerce", "food delivery"],
        "description_keywords": ["xgboost", "predictive model", "classification", "scikit-learn", "pipeline", "deployment", "inference", "feature store"]
    },
    "nlp_specialist": {
        "title": "NLP & LLM Specialist",
        "must_have_skills": {
            "nlp", "natural language processing", "llm", "llms", "fine-tuning",
            "lora", "qlora", "peft", "hugging face", "transformers", "text generation",
            "llama", "gpt", "bert", "ner", "named entity recognition", "tokenization"
        },
        "nice_to_have_skills": {
            "langchain", "llamaindex", "prompt engineering", "vector search", "semantic search",
            "embeddings", "rag", "pytorch"
        },
        "irrelevant_skills": {
            "computer vision", "image classification", "robotics", "speech recognition", "yolo", "opencv"
        },
        "title_keywords": ["nlp engineer", "nlp specialist", "llm engineer", "ai engineer", "nlp", "text processing"],
        "min_yoe": 3,
        "max_yoe": 8,
        "optimal_yoe_low": 4,
        "optimal_yoe_high": 6,
        "industry_keywords": ["software", "fintech"],
        "description_keywords": ["llm", "fine-tuning", "hugging face", "transformers", "text parsing", "sentiment analysis", "gpt", "rag"]
    },
    "search_retrieval_engineer": {
        "title": "Search & Information Retrieval Engineer",
        "must_have_skills": {
            "information retrieval", "elasticsearch", "opensearch", "solr", "lucene",
            "vector search", "hybrid search", "bm25", "semantic search", "indexing",
            "query parsing", "pinecone", "milvus", "qdrant", "weaviate", "faiss"
        },
        "nice_to_have_skills": {
            "embeddings", "learning to rank", "learning-to-rank", "ndcg", "mrr",
            "map", "evaluation frameworks", "recsys", "recommendation systems"
        },
        "irrelevant_skills": {
            "speech recognition", "image classification", "robotics", "computer vision", "gans", "yolo"
        },
        "title_keywords": ["search engineer", "retrieval engineer", "information retrieval engineer", "search backend", "indexing engineer"],
        "min_yoe": 4,
        "max_yoe": 9,
        "optimal_yoe_low": 5,
        "optimal_yoe_high": 7,
        "industry_keywords": ["software", "e-commerce", "fintech"],
        "description_keywords": ["elasticsearch", "indexing", "search query", "retrieval", "hybrid search", "faiss", "bm25", "recommender"]
    }
}

IT_SERVICES_COMPANIES = {
    "wipro", "tcs", "tata consultancy", "infosys", "accenture", "cognizant",
    "capgemini", "tech mahindra", "hcl", "mphasis", "l&t", "lnt", "mindtree",
    "genpact", "ibm", "ntt data", "dxc", "reliance jio"
}

PRODUCT_INDUSTRIES = {"software", "fintech", "e-commerce", "food delivery"}

GENERAL_TECH_TITLES = [
    "software engineer", "backend engineer", "full stack developer", "technical lead",
    "architect", "systems engineer", "data engineer"
]

IRRELEVANT_TITLES = [
    "marketing manager", "hr manager", "operations manager", "mechanical engineer",
    "graphic designer", "accountant", "sales executive", "customer support",
    "civil engineer", "ui/ux designer"
]


def is_honeypot(candidate):
    # Check 1: Expert proficiency skill with 0 duration used
    for s in candidate.get("skills", []):
        if s.get("proficiency") == "expert" and s.get("duration_months", 0) == 0:
            return True
            
    # Check 2: Job claimed duration exceeds actual calendar dates limit
    for job in candidate.get("career_history", []):
        start_str = job.get("start_date")
        end_str = job.get("end_date")
        dur = job.get("duration_months", 0)
        if start_str:
            try:
                start_date = datetime.strptime(start_str, "%Y-%m-%d")
                if end_str:
                    end_date = datetime.strptime(end_str, "%Y-%m-%d")
                    if end_date < start_date:
                        return True
                else:
                    end_date = datetime(2026, 6, 26)
                
                cal_months = (end_date.year - start_date.year) * 12 + (end_date.month - start_date.month)
                if dur > cal_months + 2:
                    return True
            except:
                pass
                
    # Check 3: last_active_date < signup_date
    signals = candidate.get("redrob_signals", {})
    signup_str = signals.get("signup_date")
    active_str = signals.get("last_active_date")
    if signup_str and active_str:
        try:
            signup_date = datetime.strptime(signup_str, "%Y-%m-%d")
            active_date = datetime.strptime(active_str, "%Y-%m-%d")
            if active_date < signup_date:
                return True
        except:
            pass
            
    # Check 4: Education dates
    for edu in candidate.get("education", []):
        start = edu.get("start_year")
        end = edu.get("end_year")
        if start and end and end < start:
            return True
            
    return False


def score_candidate(candidate, role_id="senior_ai_engineer", weights=None):
    if weights is None:
        weights = {
            "experience_weight": 1.0,
            "skills_weight": 1.0,
            "location_weight": 1.0,
            "behavioral_weight": 1.0,
            "consulting_filter_enabled": True
        }
        
    config = ROLE_CONFIGS.get(role_id, ROLE_CONFIGS["senior_ai_engineer"])
    
    profile = candidate.get("profile", {})
    skills = candidate.get("skills", [])
    career = candidate.get("career_history", [])
    signals = candidate.get("redrob_signals", {})
    
    # 1. Experience Score (Tailored to role YOE preferences)
    yoe = profile.get("years_of_experience", 0)
    exp_score = 0
    
    min_y = config["min_yoe"]
    max_y = config["max_yoe"]
    opt_low = config["optimal_yoe_low"]
    opt_high = config["optimal_yoe_high"]
    
    if opt_low <= yoe <= opt_high:
        exp_score = 25
    elif min_y <= yoe < opt_low or opt_high < yoe <= max_y:
        exp_score = 15
    elif (min_y - 1) <= yoe < min_y or max_y < yoe <= (max_y + 2):
        exp_score = 5
    else:
        exp_score = -20
        
    exp_score *= weights.get("experience_weight", 1.0)
        
    # 2. Location Alignment
    loc = profile.get("location", "").lower()
    country = profile.get("country", "").lower()
    willing_relocate = signals.get("willing_to_relocate", False)
    
    loc_score = 0
    if country == "india":
        if "noida" in loc or "pune" in loc:
            loc_score = 20
        elif willing_relocate:
            major_hubs = ["hyderabad", "mumbai", "delhi", "bangalore", "chennai", "indore", "jaipur", "ahmedabad", "kolkata", "kochi", "trivandrum", "gurgaon"]
            if any(hub in loc for hub in major_hubs):
                loc_score = 15
            else:
                loc_score = 10
        else:
            loc_score = -5
    else:
        loc_score = -40
        
    loc_score *= weights.get("location_weight", 1.0)
        
    # 3. Technical Skills
    skill_score = 0
    has_target_nlp_ir = False
    has_irrelevant_skills = False
    
    prof_weights = {
        "beginner": 0.2,
        "intermediate": 0.5,
        "advanced": 0.8,
        "expert": 1.0
    }
    
    matched_must_have = []
    matched_nice_to_have = []
    
    for s in skills:
        name = s.get("name", "").lower()
        prof = s.get("proficiency", "beginner")
        weight = prof_weights.get(prof, 0.2)
        dur = s.get("duration_months", 0)
        endorsements = s.get("endorsements", 0)
        
        trust_mult = (dur / 12.0) * (1.0 + endorsements / 10.0)
        trust_mult = min(max(trust_mult, 0.5), 4.0)
        skill_val = weight * trust_mult
        
        # Match Must Have
        if any(ms in name for ms in config["must_have_skills"]):
            skill_score += skill_val * 6.0
            matched_must_have.append(name)
            has_target_nlp_ir = True
        # Match Nice to Have
        elif any(ns in name for ns in config["nice_to_have_skills"]):
            skill_score += skill_val * 3.0
            matched_nice_to_have.append(name)
        # Match Irrelevant
        elif any(is_s in name for is_s in config["irrelevant_skills"]):
            has_irrelevant_skills = True
            
    if has_irrelevant_skills and not has_target_nlp_ir:
        skill_score -= 15
        
    skill_score *= weights.get("skills_weight", 1.0)
        
    # 4. Career and Title
    career_score = 0
    current_title = profile.get("current_title", "").lower()
    
    # Title matching
    if any(kw in current_title for kw in config["title_keywords"]):
        career_score += 15
    elif any(kw in current_title for kw in GENERAL_TECH_TITLES):
        career_score += 8
    elif any(kw in current_title for kw in IRRELEVANT_TITLES):
        career_score -= 40
        
    is_architect_manager = any(kw in current_title for kw in ["architect", "manager", "director", "lead"])
    if is_architect_manager:
        current_job = [j for j in career if j.get("is_current")]
        if current_job and current_job[0].get("duration_months", 0) > 18:
            career_score -= 10
            
    # Scan past descriptions for actual system implementation keywords
    relevant_exp_months = 0
    desc_keyword_matches = 0
    desc_keywords = config["description_keywords"]
    
    for job in career:
        title = job.get("title", "").lower()
        desc = job.get("description", "").lower()
        dur = job.get("duration_months", 0)
        
        is_relevant_role = any(kw in title for kw in config["title_keywords"]) or any(kw in desc for kw in desc_keywords)
        if is_relevant_role:
            relevant_exp_months += dur
            
        for kw in desc_keywords:
            if kw in desc:
                desc_keyword_matches += 1
                
    career_score += min((relevant_exp_months / 12.0) * 3.0, 15.0)
    career_score += min(desc_keyword_matches * 1.5, 12.0)
    
    # 5. Company Type and Job Stability
    company_score = 0
    total_jobs = len(career)
    it_services_jobs = 0
    product_company_jobs = 0
    total_months = 0
    
    for job in career:
        comp = job.get("company", "").lower()
        ind = job.get("industry", "").lower()
        size = job.get("company_size", "")
        dur = job.get("duration_months", 0)
        total_months += dur
        
        is_it_serv = (ind in ["it services", "consulting"]) or any(it in comp for it in IT_SERVICES_COMPANIES)
        if is_it_serv:
            it_services_jobs += 1
        else:
            if ind in PRODUCT_INDUSTRIES or size in ["11-50", "51-200", "201-500"]:
                product_company_jobs += 1
                
    if weights.get("consulting_filter_enabled", True):
        if total_jobs > 0 and it_services_jobs == total_jobs:
            company_score -= 30
            
    if product_company_jobs > 0:
        company_score += 10
        
    if total_jobs > 0:
        avg_dur = total_months / total_jobs
        if avg_dur < 18:
            company_score -= 10
        elif avg_dur >= 36:
            company_score += 5
            
    # 6. Behavioral Multipliers
    active_date_str = signals.get("last_active_date")
    resp_rate = signals.get("recruiter_response_rate", 0.0)
    notice = signals.get("notice_period_days", 30)
    open_work = signals.get("open_to_work_flag", False)
    int_rate = signals.get("interview_completion_rate", 0.0)
    offer_rate = signals.get("offer_acceptance_rate", -1.0)
    
    behavioral_mult = 1.0
    
    if active_date_str:
        try:
            active_date = datetime.strptime(active_date_str, "%Y-%m-%d")
            days_since = (datetime(2026, 6, 26) - active_date).days
            if days_since > 180:
                behavioral_mult *= (1.0 - 0.5 * weights.get("behavioral_weight", 1.0))
            elif days_since > 90:
                behavioral_mult *= (1.0 - 0.25 * weights.get("behavioral_weight", 1.0))
            elif days_since <= 30:
                behavioral_mult *= (1.0 + 0.1 * weights.get("behavioral_weight", 1.0))
        except:
            pass
            
    if resp_rate < 0.15:
        behavioral_mult *= (1.0 - 0.4 * weights.get("behavioral_weight", 1.0))
    elif resp_rate < 0.35:
        behavioral_mult *= (1.0 - 0.2 * weights.get("behavioral_weight", 1.0))
    elif resp_rate >= 0.75:
        behavioral_mult *= (1.0 + 0.1 * weights.get("behavioral_weight", 1.0))
        
    if notice <= 30:
        behavioral_mult *= (1.0 + 0.15 * weights.get("behavioral_weight", 1.0))
    elif notice == 90:
        behavioral_mult *= (1.0 - 0.2 * weights.get("behavioral_weight", 1.0))
    elif notice >= 120:
        behavioral_mult *= (1.0 - 0.5 * weights.get("behavioral_weight", 1.0))
        
    if open_work:
        behavioral_mult *= (1.0 + 0.1 * weights.get("behavioral_weight", 1.0))
    else:
        behavioral_mult *= (1.0 - 0.05 * weights.get("behavioral_weight", 1.0))
        
    if int_rate < 0.50:
        behavioral_mult *= (1.0 - 0.4 * weights.get("behavioral_weight", 1.0))
    elif int_rate >= 0.80:
        behavioral_mult *= (1.0 + 0.1 * weights.get("behavioral_weight", 1.0))
        
    if offer_rate == 0.0:
        behavioral_mult *= (1.0 - 0.4 * weights.get("behavioral_weight", 1.0))
    elif offer_rate > 0.60:
        behavioral_mult *= (1.0 + 0.1 * weights.get("behavioral_weight", 1.0))
        
    base_score = exp_score + loc_score + skill_score + career_score + company_score
    final_score = base_score * behavioral_mult
    
    subscores = {
        "exp_score": exp_score,
        "loc_score": loc_score,
        "skill_score": skill_score,
        "career_score": career_score,
        "company_score": company_score,
        "behavioral_mult": behavioral_mult,
        "matched_must_have": matched_must_have,
        "matched_nice_to_have": matched_nice_to_have,
        "yoe": yoe,
        "notice": notice,
        "location_raw": profile.get("location", ""),
        "willing_relocate": willing_relocate,
        "avg_dur": (total_months / total_jobs) if total_jobs > 0 else 0
    }
    
    return final_score, subscores


def generate_reasoning(rank, candidate, role_id, subscores):
    profile = candidate.get("profile", {})
    yoe = subscores["yoe"]
    title = profile.get("current_title", "")
    notice = subscores["notice"]
    loc = subscores["location_raw"]
    must_haves = subscores["matched_must_have"]
    nice_haves = subscores["matched_nice_to_have"]
    mult = subscores["behavioral_mult"]
    avg_dur = subscores["avg_dur"]
    
    role_name = ROLE_CONFIGS[role_id]["title"]
    
    # Analyze skills
    skills_matched = []
    if must_haves:
        skills_matched.append(f"must-have technical skills: {', '.join(must_haves[:3])}")
    if nice_haves:
        skills_matched.append(f"nice-to-have skills: {', '.join(nice_haves[:2])}")
        
    skills_sentence = ""
    if skills_matched:
        skills_sentence = " This worker possesses key " + " and ".join(skills_matched) + "."
    else:
        skills_sentence = " Focus is on adjacent tech capabilities, though direct matches are limited."
        
    # Analyze career history stability
    stability_sentence = ""
    if avg_dur > 0:
        years_per_job = avg_dur / 12.0
        if years_per_job >= 3.0:
            stability_sentence = f" They demonstrate exceptional career stability, averaging {years_per_job:.1f} years per job."
        elif years_per_job < 1.5:
            stability_sentence = f" They switch roles frequently, with an average job tenure of only {years_per_job:.1f} years."
        else:
            stability_sentence = f" They maintain a healthy tenure profile, averaging {years_per_job:.1f} years per role."

    # Analyze logistics / notice period
    logistics_sentence = ""
    if notice <= 30:
        logistics_sentence = f" Stated notice period of {notice} days is highly favorable for rapid onboarding."
    elif notice >= 90:
        logistics_sentence = f" However, their long notice period of {notice} days presents a scheduling risk."
    else:
        logistics_sentence = f" Stated notice period is standard ({notice} days)."

    # Platform signals
    signals_sentence = ""
    if mult > 1.1:
        signals_sentence = f" Highly active platform engagement (multiplier of {mult:.2f}x) indicates strong responsiveness."
    elif mult < 0.8:
        signals_sentence = f" Note: platform activity is relatively low (multiplier of {mult:.2f}x), which may slow down outreach."
    else:
        signals_sentence = f" Normal platform engagement signals detected (multiplier of {mult:.2f}x)."

    # Location alignment
    location_sentence = ""
    if "noida" in loc.lower() or "pune" in loc.lower():
        location_sentence = f" Residing in {loc} satisfies local office presence requirements."
    elif subscores["willing_relocate"]:
        location_sentence = f" Currently in {loc} but has explicitly marked willingness to relocate to Noida/Pune."
    else:
        location_sentence = f" Residing in {loc} without relocation indicators presents a location mismatch."

    # Combine into a descriptive reasoning block
    reason = f"Ranked #{rank} for the {role_name} role. Currently working as {title} with {yoe} years of overall experience.{skills_sentence}{stability_sentence}{location_sentence}{logistics_sentence}{signals_sentence}"
    return reason


def load_dataset():
    global ALL_CANDIDATES, VALID_CANDIDATES, HONEYPOT_COUNT, CANDIDATE_MAP, DATASET_STATS
    
    print(f"Loading dataset from: {CANDIDATES_PATH}...")
    start_time = datetime.now()
    
    candidates = []
    valid = []
    honeypots = 0
    lookup = {}
    
    countries = {}
    yoe_sum = 0
    skills_count = {}
    active_count = 0
    
    with open(CANDIDATES_PATH, "r", encoding="utf-8") as f:
        for line in f:
            if not line.strip():
                continue
            cand = json.loads(line)
            candidates.append(cand)
            cid = cand["candidate_id"]
            lookup[cid] = cand
            
            if is_honeypot(cand):
                honeypots += 1
            else:
                valid.append(cand)
                
                # Gather stats
                profile = cand.get("profile", {})
                country = profile.get("country", "")
                countries[country] = countries.get(country, 0) + 1
                yoe_sum += profile.get("years_of_experience", 0)
                
                for s in cand.get("skills", []):
                    sname = s.get("name", "")
                    skills_count[sname] = skills_count.get(sname, 0) + 1
                    
                active_str = cand.get("redrob_signals", {}).get("last_active_date")
                if active_str:
                    try:
                        active_date = datetime.strptime(active_str, "%Y-%m-%d")
                        days = (datetime(2026, 6, 26) - active_date).days
                        if days <= 30:
                            active_count += 1
                    except:
                        pass
                        
    ALL_CANDIDATES = candidates
    VALID_CANDIDATES = valid
    HONEYPOT_COUNT = honeypots
    CANDIDATE_MAP = lookup
    
    top_sk = sorted(skills_count.items(), key=lambda x: x[1], reverse=True)[:10]
    
    DATASET_STATS = {
        "total_pool": len(candidates),
        "valid_count": len(valid),
        "honeypots_filtered": honeypots,
        "average_experience": round(yoe_sum / len(valid), 2) if valid else 0,
        "active_last_30d": active_count,
        "countries": countries,
        "top_skills": dict(top_sk)
    }
    
    duration = (datetime.now() - start_time).total_seconds()
    print(f"Loaded {len(candidates)} candidates in {duration:.2f} seconds.")
    print(f"Filtered {honeypots} honeypots, {len(valid)} valid remain.")


def rank_candidates(role_id="senior_ai_engineer", weights=None):
    global CURRENT_RANKING
    
    start_time = datetime.now()
    scored = []
    
    for cand in VALID_CANDIDATES:
        score, subscores = score_candidate(cand, role_id, weights)
        scored.append((score, cand, subscores))
        
    scored.sort(key=lambda x: (-x[0], x[1]["candidate_id"]))
    
    CURRENT_RANKING = []
    for i, (score, cand, subscores) in enumerate(scored[:100]):
        rank = i + 1
        reason = generate_reasoning(rank, cand, role_id, subscores)
        CURRENT_RANKING.append({
            "rank": rank,
            "candidate_id": cand["candidate_id"],
            "score": round(score, 3),
            "name": cand["profile"]["anonymized_name"],
            "title": cand["profile"]["current_title"],
            "location": cand["profile"]["location"],
            "notice_period": subscores["notice"],
            "years_of_experience": cand["profile"]["years_of_experience"],
            "reasoning": reason,
            "subscores": {
                "exp_score": subscores["exp_score"],
                "loc_score": subscores["loc_score"],
                "skill_score": subscores["skill_score"],
                "career_score": subscores["career_score"],
                "company_score": subscores["company_score"],
                "behavioral_mult": round(subscores["behavioral_mult"], 2),
                "matched_must_have": subscores["matched_must_have"],
                "matched_nice_to_have": subscores["matched_nice_to_have"]
            }
        })
        
    duration = (datetime.now() - start_time).total_seconds()
    print(f"Re-ranked {role_id} in {duration:.2f} seconds.")


# API Endpoints
@app.route("/api/stats", methods=["GET"])
def get_stats():
    return jsonify(DATASET_STATS)


@app.route("/api/roles", methods=["GET"])
def get_roles():
    roles_list = []
    for r_id, config in ROLE_CONFIGS.items():
        roles_list.append({
            "id": r_id,
            "title": config["title"],
            "min_yoe": config["min_yoe"],
            "max_yoe": config["max_yoe"]
        })
    return jsonify(roles_list)


@app.route("/api/candidates", methods=["GET"])
def get_candidates():
    return jsonify({
        "count": len(CURRENT_RANKING),
        "candidates": CURRENT_RANKING
    })


@app.route("/api/candidates/<candidate_id>", methods=["GET"])
def get_candidate_detail(candidate_id):
    cand = CANDIDATE_MAP.get(candidate_id)
    if not cand:
        return jsonify({"error": "Candidate not found"}), 404
    return jsonify(cand)


@app.route("/api/rank", methods=["POST"])
def post_rank():
    global ACTIVE_ROLE
    data = request.get_json() or {}
    
    role_id = data.get("role_id", "senior_ai_engineer")
    ACTIVE_ROLE = role_id
    
    weights = {
        "experience_weight": float(data.get("experience_weight", 1.0)),
        "skills_weight": float(data.get("skills_weight", 1.0)),
        "location_weight": float(data.get("location_weight", 1.0)),
        "behavioral_weight": float(data.get("behavioral_weight", 1.0)),
        "consulting_filter_enabled": bool(data.get("consulting_filter_enabled", True))
    }
    
    rank_candidates(role_id, weights)
    return jsonify({
        "success": True,
        "active_role": ACTIVE_ROLE,
        "count": len(CURRENT_RANKING),
        "candidates": CURRENT_RANKING
    })


@app.route("/api/export", methods=["GET"])
def export_xlsx():
    """Generate and return a styled XLSX file of the current ranking."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    role_id = request.args.get("role_id", ACTIVE_ROLE)
    role_title = ROLE_CONFIGS.get(role_id, {}).get("title", role_id)

    wb = Workbook()
    ws = wb.active
    ws.title = "Ranked Candidates"

    # --- Styles ---
    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="1a1a2e", end_color="1a1a2e", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC")
    )
    alt_fill = PatternFill(start_color="f0f4ff", end_color="f0f4ff", fill_type="solid")
    score_font_high = Font(name="Calibri", bold=True, color="006100", size=10)
    score_font_mid = Font(name="Calibri", bold=False, color="9C5700", size=10)
    score_font_low = Font(name="Calibri", bold=False, color="9C0006", size=10)

    # --- Title Row ---
    ws.merge_cells("A1:N1")
    title_cell = ws["A1"]
    title_cell.value = f"AURA RECRUIT — Ranked Shortlist: {role_title}"
    title_cell.font = Font(name="Calibri", bold=True, size=14, color="1a1a2e")
    title_cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 32

    ws.merge_cells("A2:N2")
    subtitle_cell = ws["A2"]
    subtitle_cell.value = f"Generated on {datetime.now().strftime('%Y-%m-%d %H:%M')}  |  Pool: {len(VALID_CANDIDATES):,} valid candidates  |  Honeypots filtered: {HONEYPOT_COUNT:,}"
    subtitle_cell.font = Font(name="Calibri", size=9, italic=True, color="666666")
    ws.row_dimensions[2].height = 20

    # --- Headers (Row 4) ---
    headers = [
        "Rank", "Candidate ID", "Name", "Current Title",
        "YOE", "Location", "Notice (Days)",
        "Total Score", "Skills Score", "Experience Score",
        "Career Score", "Company Score", "Location Score",
        "Behavioral Multiplier", "Must-Have Skills Matched",
        "Nice-to-Have Skills Matched", "AI Suitability Analysis"
    ]
    col_widths = [6, 16, 22, 28, 6, 20, 12, 12, 12, 14, 12, 12, 12, 16, 35, 30, 60]

    for col_idx, (header, width) in enumerate(zip(headers, col_widths), 1):
        cell = ws.cell(row=4, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[4].height = 28

    # --- Data Rows ---
    for row_idx, cand in enumerate(CURRENT_RANKING, start=5):
        sub = cand.get("subscores", {})
        must_skills = ", ".join(sub.get("matched_must_have", []))
        nice_skills = ", ".join(sub.get("matched_nice_to_have", []))

        row_data = [
            cand["rank"],
            cand["candidate_id"],
            cand["name"],
            cand["title"],
            cand["years_of_experience"],
            cand["location"],
            cand["notice_period"],
            cand["score"],
            sub.get("skill_score", 0),
            sub.get("exp_score", 0),
            sub.get("career_score", 0),
            sub.get("company_score", 0),
            sub.get("loc_score", 0),
            round(sub.get("behavioral_mult", 1.0), 2),
            must_skills,
            nice_skills,
            cand.get("reasoning", "")
        ]

        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = thin_border
            cell.alignment = Alignment(vertical="top", wrap_text=(col_idx >= 15))
            cell.font = Font(name="Calibri", size=10)

            # Alternating row shading
            if (row_idx - 5) % 2 == 1:
                cell.fill = alt_fill

        # Color-code the total score
        score_cell = ws.cell(row=row_idx, column=8)
        score_val = cand["score"]
        if score_val >= 200:
            score_cell.font = score_font_high
        elif score_val >= 100:
            score_cell.font = score_font_mid
        else:
            score_cell.font = score_font_low

    # --- Freeze panes (freeze header row + rank column) ---
    ws.freeze_panes = "B5"

    # --- Auto-filter ---
    ws.auto_filter.ref = f"A4:Q{4 + len(CURRENT_RANKING)}"

    # --- Write to buffer ---
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"aura_recruit_{role_id}_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename
    )


@app.route("/", methods=["GET"])
def serve_index():
    if getattr(sys, 'frozen', False):
        static_dir = sys._MEIPASS
    else:
        static_dir = os.path.dirname(os.path.abspath(__file__))
    return send_from_directory(static_dir, "index.html")


if __name__ == "__main__":
    load_dataset()
    rank_candidates("senior_ai_engineer")
    
    is_frozen = getattr(sys, 'frozen', False)
    
    if is_frozen:
        # Running as a compiled EXE — open browser automatically after a short delay
        import threading
        import webbrowser
        def open_browser():
            import time
            time.sleep(2.5)
            webbrowser.open("http://127.0.0.1:5000")
        threading.Thread(target=open_browser, daemon=True).start()
        
        print("\n" + "="*60)
        print("  AURA RECRUIT — AI Candidate Ranker")
        print("  Server running at: http://127.0.0.1:5000")
        print("  Opening browser automatically...")
        print("  To stop: close this window or press Ctrl+C")
        print("="*60 + "\n")
        
        # Disable debug + use_reloader in EXE mode (reloader breaks PyInstaller)
        app.run(host="127.0.0.1", port=5000, debug=False, use_reloader=False)
    else:
        # Development mode — Flask debug with hot reload
        app.run(host="127.0.0.1", port=5000, debug=True)
